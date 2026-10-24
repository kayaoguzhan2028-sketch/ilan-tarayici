# -*- coding: utf-8 -*-
"""Tarayıcıların ortak parçaları: Google Drive'a Excel kaydı + telefona bildirim."""
import os, string, datetime, urllib.request
from pathlib import Path

# GitHub Actions'ta mı çalışıyoruz? Orada çıktılar repodaki "sonuclar" klasörüne gider.
ACTIONS = bool(os.environ.get("GITHUB_ACTIONS"))
CIKTI = Path(__file__).resolve().parent / ("sonuclar" if ACTIONS else "")
CIKTI.mkdir(exist_ok=True)
EXCEL_YERI = "GitHub > ilan-tarayici > sonuclar" if ACTIONS else "Drive > İlan Taramaları"

# Telefona bildirim: ntfy uygulamasında bu konuya abone ol (hesap/şifre gerekmez).
# Public repoda konu adı görünmesin: GitHub Secrets > NTFY_KONU
NTFY_KONU = os.environ.get("NTFY_KONU", "")
DRIVE_ALT_KLASOR = "İlan Taramaları"


def drive_klasoru():
    """Google Drive masaüstü uygulamasının 'Drive'ım' klasörünü bulur (G:, H:, I: ...)."""
    if ACTIONS:
        return None
    for harf in string.ascii_uppercase[3:]:
        for ad in ("Drive'ım", "My Drive", "Drive’ım"):
            yol = Path(f"{harf}:/") / ad
            try:
                if yol.is_dir():
                    hedef = yol / DRIVE_ALT_KLASOR
                    hedef.mkdir(exist_ok=True)
                    return hedef
            except OSError:
                continue
    return None


def excel_kaydet(dosya_adi, basliklar, satirlar, link_sutunu=None, yeni_sutunu=None):
    """satirlar: list of list. Hem yerel klasöre hem Google Drive'a kaydeder. Kaydedilen yolları döner."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.append(basliklar)
    for h in ws[1]:
        h.font = Font(bold=True)
    sari = PatternFill("solid", fgColor="FFF2B3")
    for satir in satirlar:
        ws.append(satir)
        r = ws.max_row
        if link_sutunu is not None and satir[link_sutunu]:
            c = ws.cell(row=r, column=link_sutunu + 1)
            c.hyperlink = satir[link_sutunu]
            c.font = Font(color="0563C1", underline="single")
        if yeni_sutunu is not None and satir[yeni_sutunu]:
            for c in ws[r]:
                c.fill = sari
    for kol in ws.columns:
        en = max(len(str(c.value or "")) for c in kol)
        ws.column_dimensions[kol[0].column_letter].width = min(max(10, en + 2), 60)
    ws.freeze_panes = "A2"

    # CSV kopyası: Claude'un günlük görevi Drive'dan bunu okuyup siteye (Supabase) işler
    import csv, io
    tampon = io.StringIO()
    csv.writer(tampon).writerows([basliklar] + satirlar)
    csv_adi = dosya_adi.rsplit(".", 1)[0] + ".csv"

    kaydedilen = []
    klasorler = [CIKTI]
    d = drive_klasoru()
    if d:
        klasorler.append(d)
    for k in klasorler:
        wb.save(k / dosya_adi)
        (k / csv_adi).write_text(tampon.getvalue(), encoding="utf-8-sig")
        kaydedilen.append(str(k / dosya_adi))
    return kaydedilen


def bildir(baslik, mesaj, oncelik="default"):
    """Telefona ntfy bildirimi gönderir. Hata olursa sessizce geçer."""
    if not NTFY_KONU:
        print("⚠ NTFY_KONU tanımlı değil, bildirim atlandı.")
        return False
    try:
        req = urllib.request.Request(
            f"https://ntfy.sh/{NTFY_KONU}", data=mesaj.encode("utf-8"),
            headers={"Title": baslik.translate(str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")),
                     "Priority": oncelik, "Tags": "office"})
        urllib.request.urlopen(req, timeout=15)
        return True
    except Exception as e:
        print("⚠ Bildirim gönderilemedi:", e)
        return False


def bugun():
    return datetime.date.today().isoformat()
